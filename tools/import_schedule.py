from __future__ import annotations

import json
import shutil
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
DEFAULT_SOURCE = Path(r"C:\Users\Bemister\Downloads\War Days and Caps.xlsx")


def clean_key(value: object) -> str:
    return "".join(ch for ch in str(value or "").lower() if ch.isalnum())


def normalize_date(value: object) -> str:
    if isinstance(value, datetime):
      return value.date().isoformat()
    if isinstance(value, date):
      return value.isoformat()
    text = str(value or "").strip()
    if not text:
      return ""
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%A", "%a"):
      try:
        parsed = datetime.strptime(text, fmt)
        if fmt in ("%A", "%a"):
          return text.lower()
        return parsed.date().isoformat()
      except ValueError:
        pass
    return text.lower()


def find_value(row: dict[str, object], names: tuple[str, ...]) -> object:
    for key, value in row.items():
        cleaned = clean_key(key)
        if any(name in cleaned for name in names):
            return value
    return ""


def normalize_tier(value: object) -> str:
    text = str(value or "").strip()
    lower = text.lower()
    if "t1" in lower:
        return "T1 Capped"
    if "t2" in lower:
        return "T2 Capped"
    if "uncap" in lower:
        return "Uncapped"
    return text


def rows_from_sheet(ws) -> list[dict[str, object]]:
    values = list(ws.iter_rows(values_only=True))
    best_rows: list[dict[str, object]] = []
    for header_index, header in enumerate(values[:10]):
        headers = [str(cell or "").strip() for cell in header]
        if not any(headers):
            continue
        rows = []
        for raw in values[header_index + 1 :]:
            row = {headers[i]: raw[i] if i < len(raw) else "" for i in range(len(headers))}
            node = find_value(row, ("node", "territory", "area"))
            day = find_value(row, ("date", "day"))
            if node and day:
                rows.append(row)
        if len(rows) > len(best_rows):
            best_rows = rows
    return best_rows


def rows_from_matrix(ws) -> list[dict[str, str]]:
    values = list(ws.iter_rows(values_only=True))
    if len(values) < 3:
        return []
    first_header = str(values[0][0] or "").strip().lower()
    if "day" not in first_header:
        return []

    tier_headers: list[str] = []
    current_tier = ""
    for cell in values[0]:
        if cell:
            current_tier = normalize_tier(cell)
        tier_headers.append(current_tier)

    node_headers = [str(cell or "").strip() for cell in values[1]]
    rows: list[dict[str, str]] = []
    for raw in values[2:]:
        day = normalize_date(raw[0])
        if not day:
            continue
        for index, node in enumerate(node_headers[1:], start=1):
            if not node:
                continue
            cap_value = raw[index] if index < len(raw) else ""
            if cap_value in (None, ""):
                continue
            cap = str(int(cap_value)) if isinstance(cap_value, float) and cap_value.is_integer() else str(cap_value)
            rows.append(
                {
                    "date": day,
                    "node": node,
                    "server": node,
                    "tier": tier_headers[index] or "",
                    "cap": cap,
                }
            )
    return rows


def normalize_rows(raw_rows: list[dict[str, object]]) -> list[dict[str, str]]:
    normalized = []
    for row in raw_rows:
        day = normalize_date(find_value(row, ("date", "day")))
        node = str(find_value(row, ("node", "territory", "area")) or "").strip()
        server = str(find_value(row, ("server", "channel", "region")) or "").strip()
        tier = normalize_tier(find_value(row, ("tier", "level")))
        cap = str(find_value(row, ("cap", "member", "player", "limit")) or "").strip()
        if day and node:
            normalized.append(
                {
                    "date": day,
                    "node": node,
                    "server": server,
                    "tier": tier,
                    "cap": cap or tier,
                }
            )
    return normalized


def main() -> int:
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCE
    if not source.exists():
        print(f"Workbook not found: {source}", file=sys.stderr)
        return 1

    DATA_DIR.mkdir(exist_ok=True)
    workbook_copy = DATA_DIR / "War Days and Caps.xlsx"
    if source.resolve() != workbook_copy.resolve():
        shutil.copy2(source, workbook_copy)

    wb = openpyxl.load_workbook(source, data_only=True)
    raw_rows = []
    matrix_rows = []
    for ws in wb.worksheets:
        matrix_rows.extend(rows_from_matrix(ws))
        raw_rows.extend(rows_from_sheet(ws))

    rows = matrix_rows or normalize_rows(raw_rows)
    data_js = "var BDO_WAR_SCHEDULE = " + json.dumps(rows, indent=2, ensure_ascii=False) + ";\n"
    (DATA_DIR / "schedule-data.js").write_text(data_js, encoding="utf-8")
    (DATA_DIR / "schedule-data.json").write_text(
        json.dumps(rows, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    app_js = ROOT / "app.js"
    app_text = app_js.read_text(encoding="utf-8")
    marker = "const BUNDLED_WAR_SCHEDULE = "
    start = app_text.index(marker)
    end = app_text.index(";\n\nconst presets", start)
    app_js.write_text(
        app_text[:start] + marker + json.dumps(rows, indent=2, ensure_ascii=False) + app_text[end:],
        encoding="utf-8",
    )
    print(f"Imported {len(rows)} rows from {source}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
